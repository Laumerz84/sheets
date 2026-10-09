"""Builds finance-calculator.xlsx: the finance calculator with colors, number
formats, borders and conditional formatting (works in Sheets and in Excel)."""
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "finance-calculator.xlsx")

# ---------------------------------------------------------------- palette
GREEN_DARK = "217346"
SECTIONS = {"mortgage": "2F5597", "savings": "548235", "budget": "C55A11",
            "math": "7030A0", "big": "404040"}
LIGHT = {"mortgage": "DDEBF7", "savings": "E2EFDA", "budget": "FCE4D6",
         "math": "EDE2F6", "big": "EDEDED"}
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
INPUT_FONT = Font(name="Calibri", size=11, color="0000FF", bold=True)
RESULT_FILL = PatternFill("solid", fgColor="E2EFDA")
HEAD_FILL = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
medium = Side(style="medium", color="7F7F7F")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

MONEY = '"$"#,##0.00'
MONEY0 = '"$"#,##0'
PCT = "0.0%"

wb = Workbook()
ws = wb.active
ws.title = "Calculator"
ws.sheet_view.showGridLines = False
ws.sheet_view.zoomScale = 100


def cell(addr, value=None, fmt=None, bold=False, color=None, fill=None, size=None,
         align=None, italic=False, border=None, font=None):
    c = ws[addr]
    if value is not None:
        c.value = value
    if fmt:
        c.number_format = fmt
    c.font = font or Font(name="Calibri", size=size or 11, bold=bold, italic=italic, color=color)
    if fill:
        c.fill = fill if isinstance(fill, PatternFill) else PatternFill("solid", fgColor=fill)
    if align:
        c.alignment = Alignment(horizontal=align, vertical="center")
    if border:
        c.border = border
    return c


def section(rng, title, key):
    first = rng.split(":")[0]
    ws.merge_cells(rng)
    cell(first, title, bold=True, color="FFFFFF", fill=SECTIONS[key], size=12)
    ws[first].alignment = Alignment(horizontal="left", vertical="center", indent=1)


def label(addr, text, key=None):
    cell(addr, text, fill=LIGHT[key] if key else None)
    ws[addr].alignment = Alignment(indent=1, vertical="center")


def inp(addr, value, fmt):
    cell(addr, value, fmt=fmt, font=INPUT_FONT, fill=INPUT_FILL, border=BOX, align="right")


def out(addr, formula, fmt, strong=False):
    cell(addr, formula, fmt=fmt, bold=strong, fill=RESULT_FILL if strong else None,
         border=BOX, align="right")


def header_row(row, cols, names):
    for c, n in zip(cols, names):
        cell(f"{c}{row}", n, bold=True, fill=HEAD_FILL, align="center",
             border=Border(bottom=medium, top=thin, left=thin, right=thin))


# ---------------------------------------------------------------- title
ws.merge_cells("A1:H1")
cell("A1", "PERSONAL FINANCE CALCULATOR", bold=True, size=18, color="FFFFFF", fill=GREEN_DARK)
ws["A1"].alignment = Alignment(vertical="center", indent=1)
ws.merge_cells("J1:K1")
cell("J1", "=TODAY()", fmt="dddd, mmmm d, yyyy", bold=True, color="FFFFFF", fill=GREEN_DARK, align="right")
ws.row_dimensions[1].height = 34
cell("A2", "Yellow cells are inputs - change them and everything recalculates.", italic=True, color="7F7F7F")
cell("J2", "Input", font=INPUT_FONT, fill=INPUT_FILL, border=BOX, align="center")
cell("K2", "Result", bold=True, fill=RESULT_FILL, border=BOX, align="center")

# ---------------------------------------------------------------- mortgage
section("A4:B4", "MORTGAGE", "mortgage")
rows = [
    ("Home price", ("in", 450000, MONEY0)),
    ("Down payment", ("in", 0.20, "0%")),
    ("Interest rate (APR)", ("in", 0.065, "0.00%")),
    ("Loan term (years)", ("in", 30, "0")),
    ("Property tax / year", ("in", 5400, MONEY0)),
    ("Insurance / year", ("in", 1800, MONEY0)),
    ("Loan amount", ("out", "=B5*(1-B6)", MONEY0, False)),
    ("Monthly principal + interest", ("out", "=-PMT(B7/12,B8*12,B11)", MONEY, True)),
    ("Monthly total incl. tax + ins.", ("out", "=B12+(B9+B10)/12", MONEY, True)),
    ("Total paid over the loan", ("out", "=B12*B8*12", MONEY0, False)),
    ("Total interest", ("out", "=B14-B11", MONEY0, False)),
    ("Interest as % of loan", ("out", "=B15/B11", PCT, False)),
    ("Paid off in", ("out", "=EDATE(TODAY(),B8*12)", "mmmm yyyy", False)),
]
for r, (text, spec) in enumerate(rows, start=5):
    label(f"A{r}", text, "mortgage")
    if spec[0] == "in":
        inp(f"B{r}", spec[1], spec[2])
    else:
        out(f"B{r}", spec[1], spec[2], spec[3])

section("D4:H4", "FIRST YEAR OF PAYMENTS", "mortgage")
header_row(5, "DEFGH", ["Month", "Payment", "Interest", "Principal", "Balance"])
for r in range(6, 18):
    prev = "$B$11" if r == 6 else f"H{r - 1}"
    band = PatternFill("solid", fgColor=LIGHT["mortgage"]) if r % 2 == 0 else None
    cell(f"D{r}", 1 if r == 6 else f"=D{r - 1}+1", fmt="0", align="center", border=BOX, fill=band)
    cell(f"E{r}", "=$B$12", fmt=MONEY, border=BOX, fill=band)
    cell(f"F{r}", f"={prev}*$B$7/12", fmt=MONEY, border=BOX, fill=band)
    cell(f"G{r}", f"=E{r}-F{r}", fmt=MONEY, border=BOX, fill=band)
    cell(f"H{r}", f"={prev}-G{r}", fmt=MONEY0, border=BOX, fill=band)
cell("D18", "Year 1", bold=True, align="center", border=BOX, fill=HEAD_FILL)
for c in "EFG":
    cell(f"{c}18", f"=SUM({c}6:{c}17)", fmt=MONEY0, bold=True, border=BOX, fill=HEAD_FILL)
cell("H18", "=1-H17/B11", fmt='"Paid off "0.0%', bold=True, border=BOX, fill=HEAD_FILL, align="right")

# ---------------------------------------------------------------- savings
section("A20:B20", "SAVINGS & INVESTING", "savings")
rows = [
    ("Starting balance", ("in", 10000, MONEY0)),
    ("Monthly contribution", ("in", 750, MONEY0)),
    ("Expected annual return", ("in", 0.07, "0.0%")),
    ("Years", ("in", 25, "0")),
    ("Inflation", ("in", 0.03, "0.0%")),
    ("Future value", ("out", "=FV(B23/12,B24*12,-B22,-B21)", MONEY0, True)),
    ("Total you put in", ("out", "=B21+B22*12*B24", MONEY0, False)),
    ("Growth from returns", ("out", "=B26-B27", MONEY0, False)),
    ("Future value in today's money", ("out", "=B26/(1+B25)^B24", MONEY0, True)),
    ("Monthly income (4% rule)", ("out", "=B26*0.04/12", MONEY0, False)),
    ("Years to double (rule of 72)", ("out", "=72/(B23*100)", '0.0" yrs"', False)),
]
for r, (text, spec) in enumerate(rows, start=21):
    label(f"A{r}", text, "savings")
    if spec[0] == "in":
        inp(f"B{r}", spec[1], spec[2])
    else:
        out(f"B{r}", spec[1], spec[2], spec[3])

section("D20:H20", "GROWTH BY YEAR", "savings")
header_row(21, "DEFGH", ["Year", "Start", "Added", "Growth", "End"])
for r in range(22, 32):
    band = PatternFill("solid", fgColor=LIGHT["savings"]) if r % 2 == 1 else None
    cell(f"D{r}", 1 if r == 22 else f"=D{r - 1}+1", fmt="0", align="center", border=BOX, fill=band)
    cell(f"E{r}", "=$B$21" if r == 22 else f"=H{r - 1}", fmt=MONEY0, border=BOX, fill=band)
    cell(f"F{r}", "=$B$22*12", fmt=MONEY0, border=BOX, fill=band)
    cell(f"G{r}", f"=H{r}-E{r}-F{r}", fmt=MONEY0, border=BOX, fill=band)
    cell(f"H{r}", f"=FV($B$23/12,12,-$B$22,-E{r})", fmt=MONEY0, bold=True, border=BOX, fill=band)

# ---------------------------------------------------------------- budget
section("A34:F34", "MONTHLY BUDGET", "budget")
label("A35", "Monthly take-home pay", "budget")
inp("B35", 7800, MONEY0)
header_row(36, "ABCDEF", ["Category", "Budget", "Spent", "Left", "% of pay", "Status"])
budget = [
    ("Housing (from mortgage)", "=B13", "=B13"),
    ("Groceries", 650, 712),
    ("Transport", 400, 365),
    ("Utilities", 280, 301),
    ("Dining out", 300, 420),
    ("Savings (from above)", "=B22", "=B22"),
    ("Health", 200, 150),
    ("Fun money", 250, 180),
]
for r, (cat, b, s) in enumerate(budget, start=37):
    linked = isinstance(b, str)
    cell(f"A{r}", cat, border=BOX, italic=linked)
    ws[f"A{r}"].alignment = Alignment(indent=1)
    if linked:
        cell(f"B{r}", b, fmt=MONEY0, border=BOX, italic=True, color="595959")
        cell(f"C{r}", s, fmt=MONEY0, border=BOX, italic=True, color="595959")
    else:
        inp(f"B{r}", b, MONEY0)
        inp(f"C{r}", s, MONEY0)
    cell(f"D{r}", f"=B{r}-C{r}", fmt='"$"#,##0;[Red]-"$"#,##0', border=BOX)
    cell(f"E{r}", f"=C{r}/$B$35", fmt=PCT, border=BOX)
    cell(f"F{r}", f'=IF(C{r}>B{r},"Over by "&TEXT(C{r}-B{r},"$#,##0"),"OK")', border=BOX, align="center")
for c, f, fmt in (("A", "Total", None), ("B", "=SUM(B37:B44)", MONEY0), ("C", "=SUM(C37:C44)", MONEY0),
                  ("D", "=B45-C45", '"$"#,##0;[Red]-"$"#,##0'), ("E", "=C45/B35", PCT),
                  ("F", '=IF(C45>B35,"Spending > income!","Left: "&TEXT(B35-C45,"$#,##0"))', None)):
    cell(f"{c}45", f, fmt=fmt, bold=True, fill=LIGHT["budget"],
         border=Border(top=medium, bottom=medium, left=thin, right=thin), align="center" if c == "F" else None)
extras = [
    ("Biggest expense", "=INDEX(A37:A44,MATCH(MAX(C37:C44),C37:C44,0))", None),
    ("Categories over budget", '=COUNTIF(D37:D44,"<0")', "0"),
    ("Average spend per category", "=AVERAGE(C37:C44)", MONEY),
    ("Savings rate", "=C42/B35", PCT),
    ("Spend excl. housing + savings", '=SUMIFS(C37:C44,A37:A44,"<>Housing*",A37:A44,"<>Savings*")', MONEY0),
]
for r, (text, f, fmt) in enumerate(extras, start=46):
    label(f"A{r}", text, "budget")
    out(f"B{r}", f, fmt)

# ---------------------------------------------------------------- quick math
section("J4:K4", "QUICK MATH", "math")
label("J5", "Any number", "math")
inp("K5", 144, "#,##0")
math_rows = [
    ("Square root", "=SQRT(K5)", "General"),
    ("Squared", "=K5^2", "#,##0"),
    ("Even?", '=IF(ISEVEN(K5),"yes","no")', None),
    ("As % of 1,000", "=K5/1000", PCT),
    ("Factorial of 10", "=FACT(10)", "#,##0"),
    ("Pi", "=PI()", "0.000000"),
    ("Random 1-100 (F9 rerolls)", "=RANDBETWEEN(1,100)", "0"),
    ("Days until New Year", "=DATE(YEAR(TODAY())+1,1,1)-TODAY()", '0" days"'),
    ("Week of the year", "=WEEKNUM(TODAY())", "0"),
    ("18% tip on $86.40", "=86.4*0.18", MONEY),
    ("...total split 4 ways", "=86.4*1.18/4", MONEY),
    ("22 C in Fahrenheit", "=22*9/5+32", '0.0" F"'),
    ("Marathon in km", "=26.2*1.609344", '0.00" km"'),
]
for r, (text, f, fmt) in enumerate(math_rows, start=6):
    label(f"J{r}", text, "math")
    out(f"K{r}", f, fmt)

# ---------------------------------------------------------------- big picture
section("J20:K20", "THE BIG PICTURE", "big")
big = [
    ("Housing cost as % of pay", "=B13/B35", PCT, True),
    ("vs. the 28% guideline", '=IF(B13/B35>0.28,"Above","Within")', None, False),
    ("Savings in 25 yrs (today's $)", "=B29", MONEY0, True),
    ("Years of spending it covers", "=B29/(C45*12)", '0.0" yrs"', False),
    ("Mortgage interest per $1 borrowed", "=B15/B11", '"$"0.00', False),
    ("Net cash flow / year", "=(B35-C45)*12", '"$"#,##0;[Red]-"$"#,##0', True),
]
for r, (text, f, fmt, strong) in enumerate(big, start=21):
    label(f"J{r}", text, "big")
    out(f"K{r}", f, fmt, strong)

# ---------------------------------------------------------------- conditional formatting
ws.conditional_formatting.add("F37:F44", FormulaRule(formula=['C37>B37'], fill=PatternFill(bgColor="FFC7CE"),
                                                     font=Font(color="9C0006", bold=True)))
ws.conditional_formatting.add("F37:F44", FormulaRule(formula=['C37<=B37'], fill=PatternFill(bgColor="C6EFCE"),
                                                     font=Font(color="006100")))
ws.conditional_formatting.add("E37:E44", DataBarRule(start_type="num", start_value=0, end_type="max",
                                                     color="F4B183"))
ws.conditional_formatting.add("F6:F17", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                       end_type="max", end_color="F8CBAD"))
ws.conditional_formatting.add("H22:H31", DataBarRule(start_type="num", start_value=0, end_type="max",
                                                     color="70AD47"))
ws.conditional_formatting.add("K22", FormulaRule(formula=['K21>0.28'], fill=PatternFill(bgColor="FFC7CE"),
                                                 font=Font(color="9C0006", bold=True)))
ws.conditional_formatting.add("K22", FormulaRule(formula=['K21<=0.28'], fill=PatternFill(bgColor="C6EFCE"),
                                                 font=Font(color="006100", bold=True)))
ws.conditional_formatting.add("K26", CellIsRule(operator="lessThan", formula=["0"],
                                                fill=PatternFill(bgColor="FFC7CE")))

# ---------------------------------------------------------------- layout
for col, width in {"A": 33, "B": 15, "C": 13, "D": 12, "E": 13, "F": 15, "G": 13, "H": 15,
                   "I": 3, "J": 32, "K": 17}.items():
    ws.column_dimensions[col].width = width
for r in (4, 20, 34):
    ws.row_dimensions[r].height = 22
ws.freeze_panes = "A3"

wb.save(OUT)
print(OUT)
