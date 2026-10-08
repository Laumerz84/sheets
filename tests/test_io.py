import os

import pytest

from sheets import errors, ops
from sheets.fileio import open_file, save_file
from sheets.refs import key, parse_addr
from sheets.workbook import new_workbook


def put(sh, a, text):
    r, c = parse_addr(a)
    sh.set_state(key(r, c), ops.input_state(sh, r, c, text))


def test_csv_roundtrip_preserves_text(tmp_path):
    src = tmp_path / "data.csv"
    content = "id,zip,price,date,note\r\n1,00501,1.50,2024-03-05,\"has, comma\"\r\n2,90210,\"$1,234.00\",3/5/2024,\"multi\nline\"\r\n"
    src.write_bytes(content.encode("utf-8"))
    wb = open_file(str(src))
    sh = wb.sheets[0]
    assert sh.value(1, 1) == "00501"
    assert sh.value(1, 2) == 1.5
    assert sh.value(1, 3) == 45356
    assert sh.value(2, 4) == "multi\nline"
    out = tmp_path / "out.csv"
    save_file(wb, str(out))
    assert out.read_bytes().decode("utf-8") == content


def test_csv_semicolon_and_cp1252(tmp_path):
    src = tmp_path / "eu.csv"
    src.write_bytes("name;city\r\nJosé;Málaga\r\n".encode("cp1252"))
    wb = open_file(str(src))
    sh = wb.sheets[0]
    assert sh.value(1, 1) == "Málaga"
    save_file(wb, str(src))
    assert src.read_bytes().decode("cp1252") == "name;city\r\nJosé;Málaga\r\n"


def test_xlsx_roundtrip(tmp_path):
    wb = new_workbook()
    sh = wb.sheets[0]
    put(sh, "A1", "Item")
    put(sh, "B1", "Cost")
    put(sh, "A2", "Widget")
    put(sh, "B2", "$1,200.50")
    put(sh, "A3", "Gadget")
    put(sh, "B3", "99")
    put(sh, "B4", "=SUM(B2:B3)")
    put(sh, "C1", "2024-03-05")
    put(sh, "D1", "=XLOOKUP(\"Gadget\",A2:A3,B2:B3)")
    wb.recalc()
    st = sh.style(0, 0).with_(bold=True, fill="#FFFF00")
    sh.set_style(0, 0, st)
    sh.col_widths[0] = 150
    sh.merges.append((5, 0, 5, 2))
    sh.freeze = (1, 0)
    s2 = wb.add_sheet("Other")
    s2.set_formula(0, 0, "=Sheet1!B4*2")
    wb.recalc()
    path = tmp_path / "book.xlsx"
    save_file(wb, str(path))

    wb2 = open_file(str(path))
    a, b = wb2.sheets
    assert a.value(1, 1) == 1200.5
    assert a.formula_text(3, 1) == "=SUM(B2:B3)"
    assert a.value(3, 1) == 1299.5
    assert a.value(0, 3) == 99
    assert a.style(0, 0).bold and a.style(0, 0).fill == "#FFFF00"
    assert a.style(1, 1).numfmt == "$#,##0.00"
    assert a.style(0, 2).numfmt == "yyyy-mm-dd"
    assert abs(a.col_widths[0] - 150) <= 2
    assert (5, 0, 5, 2) in a.merges
    assert a.freeze == (1, 0)
    assert b.value(0, 0) == 2599
    # check Excel will see the new function properly prefixed
    import zipfile
    with zipfile.ZipFile(path) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode()
    assert "_xlfn.XLOOKUP" in xml

    # edit and save over the same file (reusing the loaded workbook)
    put(a, "B3", "1")
    wb2.recalc()
    a.insert_rows(0, 1)
    save_file(wb2, str(path))
    wb3 = open_file(str(path))
    assert wb3.sheets[0].value(4, 1) == 1201.5
    assert wb3.sheets[1].formula_text(0, 0) == "=Sheet1!B5*2"
    assert wb3.sheets[0].style(1, 0).bold


def test_xls_open(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    path = tmp_path / "old.xls"
    book = xlwt.Workbook()
    ws = book.add_sheet("S")
    ws.write(0, 0, "Hello")
    ws.write(1, 0, 42)
    book.save(str(path))
    wb = open_file(str(path))
    assert wb.sheets[0].value(0, 0) == "Hello"
    assert wb.sheets[0].value(1, 0) == 42


def test_xlsx_preserves_unmodelled_features(tmp_path):
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import PatternFill
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
    p = tmp_path / "rich.xlsx"
    b = openpyxl.Workbook()
    ws = b.active
    ws.title = "Data"
    for i in range(1, 6):
        ws.cell(i, 1, i * 10)
    ws["B1"] = "=SUM(Prices)"
    b.defined_names["Prices"] = DefinedName("Prices", attr_text="Data!$A$1:$A$5")
    ws.conditional_formatting.add("A1:A5", CellIsRule(operator="greaterThan", formula=["25"],
                                                      fill=PatternFill("solid", fgColor="FFC7CE")))
    dv = DataValidation(type="list", formula1='"yes,no"')
    ws.add_data_validation(dv)
    dv.add("C1:C5")
    ws["D1"].comment = Comment("a note", "me")
    ws["E1"].fill = PatternFill("solid", fgColor=openpyxl.styles.colors.Color(theme=4, tint=0.6))
    b.save(p)

    wb = open_file(str(p))
    sh = wb.sheets[0]
    assert sh.value(0, 1) == 150  # defined name evaluated
    assert sh.style(0, 4).fill is not None
    put(sh, "A6", "60")
    wb.recalc()
    save_file(wb, str(p))

    b2 = openpyxl.load_workbook(p)
    ws2 = b2["Data"]
    assert ws2["A6"].value == 60
    assert len(ws2.conditional_formatting) == 1
    assert len(ws2.data_validations.dataValidation) == 1
    assert ws2["D1"].comment is not None and ws2["D1"].comment.text.startswith("a note")
    assert "Prices" in b2.defined_names
    assert ws2["E1"].fill.fgColor.theme == 4  # original style kept exactly


def test_conditional_formatting_render(tmp_path):
    import openpyxl
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
    from openpyxl.styles import Font, PatternFill
    from sheets.condfmt import CFEngine
    p = tmp_path / "cf.xlsx"
    b = openpyxl.Workbook()
    ws = b.active
    for i in range(1, 11):
        ws.cell(i, 1, i)
        ws.cell(i, 2, i * 2)
    ws.conditional_formatting.add("A1:A10", CellIsRule(operator="greaterThan", formula=["5"],
                                  fill=PatternFill(bgColor="FFC7CE"), font=Font(color="9C0006")))
    ws.conditional_formatting.add("B1:B10", FormulaRule(formula=["MOD(B1,4)=0"], fill=PatternFill(bgColor="C6EFCE")))
    ws.conditional_formatting.add("C1:C10", ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                           end_type="max", end_color="FF0000"))
    for i in range(1, 11):
        ws.cell(i, 3, i)
    b.save(p)
    wb = open_file(str(p))
    sh = wb.sheets[0]
    eng = CFEngine(sh)
    assert eng.style_for(sh.cond_formats, 0, 0) is None
    assert eng.style_for(sh.cond_formats, 6, 0)["fill"] == "#FFC7CE"
    assert eng.style_for(sh.cond_formats, 6, 0)["color"] == "#9C0006"
    assert eng.style_for(sh.cond_formats, 1, 1)["fill"] == "#C6EFCE"   # B2 = 4
    assert eng.style_for(sh.cond_formats, 0, 1) is None                 # B1 = 2
    assert eng.style_for(sh.cond_formats, 9, 2)["fill"] == "#FF0000"
    assert eng.style_for(sh.cond_formats, 0, 2)["fill"] == "#FFFFFF"


def test_xls_impostors(tmp_path):
    html = tmp_path / "bank.xls"
    html.write_text("<html><body><table><tr><th>Date</th><th>Amount</th></tr>"
                    "<tr><td>2024-03-05</td><td>1,234.50</td></tr></table></body></html>", encoding="utf-8")
    wb = open_file(str(html))
    sh = wb.sheets[0]
    assert sh.value(0, 0) == "Date" and sh.style(0, 0).bold
    assert sh.value(1, 1) == 1234.5
    xl = tmp_path / "renamed.xls"
    b = new_workbook()
    b.sheets[0].set_value(0, 0, 9.0)
    save_file(b, str(tmp_path / "real.xlsx"))
    xl.write_bytes((tmp_path / "real.xlsx").read_bytes())
    assert open_file(str(xl)).sheets[0].value(0, 0) == 9


def test_structural_edits_move_xlsx_extras(tmp_path):
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
    p = tmp_path / "x.xlsx"
    b = openpyxl.Workbook()
    ws = b.active
    ws.title = "Sheet1"
    for i in range(1, 6):
        ws.cell(i, 1, i)
    ws["A3"].comment = Comment("note", "me")
    ws["B3"] = "=SUM(DATA)"
    b.defined_names["DATA"] = DefinedName("DATA", attr_text="Sheet1!$A$1:$A$5")
    dv = DataValidation(type="list", formula1='"y,n"')
    ws.add_data_validation(dv)
    dv.add("C3")
    ws.conditional_formatting.add("A3", CellIsRule(operator="greaterThan", formula=["0"],
                                                   fill=PatternFill(bgColor="FFC7CE")))
    ws["D1"] = "rotated"
    ws["D1"].alignment = Alignment(textRotation=45)
    ws["D2"].protection = Protection(locked=False)
    b.save(p)

    wb = open_file(str(p))
    sh = wb.sheets[0]
    assert sh.value(2, 1) == 15
    sh.insert_rows(0, 2)
    assert sh.value(4, 1) == 15  # name moved with the data
    assert wb.names["DATA"] == "Sheet1!$A$3:$A$7"
    wb.rename_sheet(sh, "Data Sheet")
    save_file(wb, str(p))

    b2 = openpyxl.load_workbook(p)
    ws2 = b2["Data Sheet"]
    assert ws2["A5"].comment is not None and ws2["A3"].comment is None
    assert str(ws2.data_validations.dataValidation[0].sqref) == "C5"
    assert [str(cf.sqref) for cf in ws2.conditional_formatting] == ["A5"]
    assert b2.defined_names["DATA"].attr_text == "'Data Sheet'!$A$3:$A$7"
    assert ws2["D3"].alignment.textRotation == 45     # untouched style kept exactly
    assert ws2["D4"].protection.locked is False


def test_edited_cell_keeps_workbook_default_font(tmp_path):
    import openpyxl
    from openpyxl.styles import Font
    p = tmp_path / "arial.xlsx"
    b = openpyxl.Workbook()
    b._fonts[0] = Font(name="Arial", sz=10)
    ws = b.active
    ws["A1"] = "x"
    b.save(p)
    wb = open_file(str(p))
    sh = wb.sheets[0]
    sh.set_style(0, 0, sh.style(0, 0).with_(bold=True))
    save_file(wb, str(p))
    f = openpyxl.load_workbook(p).active["A1"].font
    assert f.b and f.name == "Arial" and f.sz == 10


def test_csv_keeps_ambiguous_text(tmp_path):
    src = 'a,b,c,d,e,f,g,h,i\r\n1/2,(5),true, 42 ,2024-1-5,1.23457E+11,0.50,12.5%,5 March 2024\r\n'
    p = tmp_path / "amb.csv"
    p.write_bytes(src.encode("utf-8"))
    wb = open_file(str(p))
    sh = wb.sheets[0]
    assert sh.value(1, 0) == "1/2" and sh.value(1, 2) == "true"
    assert sh.value(1, 4) == 45296  # still a real date
    save_file(wb, str(p))
    assert p.read_bytes().decode("utf-8") == src
