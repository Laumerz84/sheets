"""Excel .xlsx/.xlsm reading and writing through openpyxl."""
import colorsys
import datetime as dt
import os
import re
import zipfile
from copy import copy

from . import errors
from .errors import XLError
from .formula import from_file_formula, to_file_formula
from .numfmt import datetime_to_serial
from .refs import col_name, parse_range, range_addr
from .workbook import (DEFAULT_COL_WIDTH, DEFAULT_ROW_HEIGHT, DEFAULT_STYLE,
                       Formula, Sheet, Style, Workbook, intern_style)

DEFAULT_THEME = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31",
                 "A5A5A5", "FFC000", "5B9BD5", "70AD47", "0563C1", "954F72"]


def _theme_colors(book):
    xml = getattr(book, "loaded_theme", None)
    if not xml:
        return DEFAULT_THEME
    if isinstance(xml, bytes):
        xml = xml.decode("utf-8", "replace")
    found = {}
    for tag in ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4",
                "accent5", "accent6", "hlink", "folHlink"):
        m = re.search(rf"<a:{tag}>(.*?)</a:{tag}>", xml, re.S)
        if m:
            c = re.search(r'(?:srgbClr val|lastClr)="([0-9A-Fa-f]{6})"', m.group(1))
            if c:
                found[tag] = c.group(1).upper()
    order = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4",
             "accent5", "accent6", "hlink", "folHlink"]
    return [found.get(t, DEFAULT_THEME[i]) for i, t in enumerate(order)]


def _apply_tint(hexrgb, tint):
    if not tint:
        return hexrgb
    r, g, b = (int(hexrgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0, min(1, l)), s)
    return "".join(f"{round(x * 255):02X}" for x in (r, g, b))


def _color(c, theme, default=None):
    """openpyxl Color -> '#RRGGBB' or default."""
    if c is None:
        return default
    try:
        t = c.type
        if t == "rgb":
            v = c.rgb
            if not isinstance(v, str) or len(v) < 6:
                return default
            if len(v) == 8 and v[:2] == "00" and v[2:] == "000000":
                return default
            return "#" + v[-6:].upper()
        if t == "theme":
            idx = c.theme
            if idx is None or idx >= len(theme):
                return default
            return "#" + _apply_tint(theme[idx], c.tint or 0)
        if t == "indexed":
            from openpyxl.styles.colors import COLOR_INDEX
            idx = c.indexed
            if idx is None or idx >= 64 or idx >= len(COLOR_INDEX):
                return default
            return "#" + COLOR_INDEX[idx][-6:]
    except (AttributeError, ValueError, TypeError):
        pass
    return default


def _detect_lost(path):
    lost = []
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
    except (zipfile.BadZipFile, OSError):
        return lost
    if any(n.startswith("xl/charts/") for n in names):
        lost.append("charts")
    if any(n.startswith("xl/media/") for n in names):
        lost.append("pictures")
    if any(n.startswith("xl/drawings/") and n.endswith(".xml") for n in names) and not lost:
        lost.append("shapes")
    if any(n.startswith("xl/slicers/") for n in names):
        lost.append("slicers")
    return lost


def _source(path):
    """openpyxl insists on an .xlsx-like extension for paths; give it bytes otherwise."""
    if os.path.splitext(path)[1].lower() in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        return path
    import io
    with open(path, "rb") as fh:
        return io.BytesIO(fh.read())


def load_xlsx(path, progress=None):
    import openpyxl
    ext = os.path.splitext(path)[1].lower()
    book = openpyxl.load_workbook(_source(path), data_only=False, keep_vba=(ext == ".xlsm"),
                                  keep_links=True, rich_text=False)
    theme = _theme_colors(book)
    wb = Workbook()
    wb.path = path
    wb.file_format = "xlsx"
    wb.xl_book = book
    wb.xl_lost_features = _detect_lost(path)
    try:
        wb.names = {}
        for name, dn in book.defined_names.items():
            if dn.attr_text and not dn.attr_text.startswith("#"):
                wb.names[name.upper()] = dn.attr_text
    except AttributeError:
        wb.names = {}

    default_font = book._fonts[0] if book._fonts else None
    base_font_name = getattr(default_font, "name", None)
    base_font_size = getattr(default_font, "sz", None)
    wb.xl_default_font = (base_font_name or "Calibri", float(base_font_size or 11))
    style_cache = {}

    def convert_style(cell):
        sa = cell._style
        k = tuple(sa)
        st = style_cache.get(k)
        if st is not None:
            return st
        font = book._fonts[sa.fontId] if sa.fontId < len(book._fonts) else None
        fill = book._fills[sa.fillId] if sa.fillId < len(book._fills) else None
        border = book._borders[sa.borderId] if sa.borderId < len(book._borders) else None
        align = book._alignments[sa.alignmentId] if sa.alignmentId < len(book._alignments) else None
        kw = {}
        if font is not None:
            kw["bold"] = bool(font.b)
            kw["italic"] = bool(font.i)
            kw["underline"] = bool(font.u) and font.u != "none"
            kw["strike"] = bool(font.strike)
            if font.name and font.name != base_font_name:
                kw["font"] = font.name
            if font.sz and font.sz != base_font_size:
                kw["size"] = float(font.sz)
            col = _color(font.color, theme)
            if col and col != "#000000":
                kw["color"] = col
        if fill is not None:
            ft = getattr(fill, "fill_type", None) or getattr(fill, "patternType", None)
            if ft and ft != "none":
                col = _color(fill.fgColor, theme) if ft == "solid" else (
                    _color(fill.fgColor, theme) or _color(fill.bgColor, theme))
                if col:
                    kw["fill"] = col
            elif getattr(fill, "stop", None):
                try:
                    kw["fill"] = _color(fill.stop[0].color, theme)
                except (IndexError, AttributeError):
                    pass
        if align is not None:
            h = align.horizontal
            if h in ("left", "center", "right"):
                kw["halign"] = h
            elif h in ("centerContinuous", "distributed"):
                kw["halign"] = "center"
            elif h == "justify":
                kw["halign"] = "left"
            v = align.vertical
            if v in ("top", "center"):
                kw["valign"] = v
            elif v in ("justify", "distributed"):
                kw["valign"] = "center"
            if align.wrap_text:
                kw["wrap"] = True
            if align.indent:
                kw["indent"] = int(align.indent)
        if border is not None:
            sides = []
            for side in (border.left, border.top, border.right, border.bottom):
                if side is not None and side.style:
                    sides.append((side.style, _color(side.color, theme, "#000000")))
                else:
                    sides.append(None)
            kw["border"] = tuple(sides)
        nf = cell.number_format
        if nf and nf != "General":
            kw["numfmt"] = nf
        st = intern_style(Style(**kw))
        style_cache[k] = st
        return st

    needs_cache = []
    controls_ws = next((ws for ws in book.worksheets if ws.title == CONTROLS_SHEET), None)
    sheets = [ws for ws in book.worksheets if ws is not controls_ws]
    for si, ws in enumerate(sheets):
        sh = Sheet(wb, ws.title)
        sh.xl = ws
        wb.sheets.append(sh)
        values, styles = sh.values, sh.styles
        xl_styles, notes = sh.xl_styles, sh.notes
        cells = ws._cells
        total = max(1, len(cells))
        for n, ((row, col), cell) in enumerate(cells.items()):
            if progress and n % 20000 == 0:
                progress((si + n / total) / len(sheets))
            r, c = row - 1, col - 1
            k = (r << 14) | c
            v = cell.value
            if cell._style is not None and any(cell._style):
                st = convert_style(cell)
                if st is not DEFAULT_STYLE:
                    styles[k] = st
                # keep the exact original style so untouched cells save unchanged
                xl_styles[k] = (st, cell._style)
            com = getattr(cell, "comment", None)
            link = getattr(cell, "_hyperlink", None)
            if com is not None or link is not None:
                notes[k] = (com, link)
            if v is None:
                continue
            dtype = cell.data_type
            if dtype == "f" or (isinstance(v, str) and v.startswith("=") and dtype != "s"):
                text = v if isinstance(v, str) else getattr(v, "text", None)
                if text is None:
                    continue
                if not text.startswith("="):
                    text = "=" + text
                f = Formula(from_file_formula(text))
                sh.formulas[k] = f
                if f.unknown or _uses_names(f):
                    needs_cache.append((sh, k))
                continue
            if hasattr(v, "text") and not isinstance(v, str):  # ArrayFormula etc.
                text = v.text if v.text.startswith("=") else "=" + v.text
                sh.formulas[k] = Formula(from_file_formula(text))
                continue
            if isinstance(v, bool):
                values[k] = v
            elif isinstance(v, (int, float)):
                values[k] = float(v)
            elif isinstance(v, (dt.datetime, dt.date, dt.time, dt.timedelta)):
                try:
                    values[k] = datetime_to_serial(v)
                except (TypeError, ValueError, OverflowError):
                    values[k] = str(v)
            elif dtype == "e" and isinstance(v, str):
                values[k] = errors.from_code(v) or v
            else:
                values[k] = str(v)

        # geometry
        for letter, cd in ws.column_dimensions.items():
            try:
                lo, hi = cd.min or 0, cd.max or 0
            except AttributeError:
                continue
            if not lo:
                continue
            for c in range(lo - 1, min(hi, lo + 300)):
                if cd.width and cd.customWidth is not False:
                    sh.col_widths[c] = max(2, int(round(cd.width * 7 + 5)))
                if cd.hidden:
                    sh.hidden_cols.add(c)
        for r, rd in ws.row_dimensions.items():
            if rd.ht:
                sh.row_heights[r - 1] = max(2, int(round(rd.ht * 96 / 72)))
            if rd.hidden:
                sh.hidden_rows.add(r - 1)
        for rng in ws.merged_cells.ranges:
            sh.merges.append((rng.min_row - 1, rng.min_col - 1, rng.max_row - 1, rng.max_col - 1))
        _read_freeze(ws, sh)
        if ws.auto_filter and ws.auto_filter.ref:
            b = parse_range(ws.auto_filter.ref)
            if b:
                sh.autofilter = b
                # rows hidden by Excel's filter come in as hidden rows
                hidden_in = {r for r in sh.hidden_rows if b[0] < r <= b[2]}
                sh.filter_hidden = hidden_in
                sh.hidden_rows -= hidden_in
        try:
            from .condfmt import load_rules
            sh.cond_formats, sh.cf_complete = load_rules(ws, lambda c: _color(c, theme))
        except Exception:
            sh.cond_formats, sh.cf_complete = [], False
        try:
            for dv in ws.data_validations.dataValidation:
                rects = [b for b in (parse_range(x) for x in str(dv.sqref).split()) if b]
                if rects:
                    sh.xl_dv.append([dv, rects])
        except AttributeError:
            pass
        try:
            tc = ws.sheet_properties.tabColor
            sh.tab_color = _color(tc, theme) if tc is not None else None
        except AttributeError:
            pass
        try:
            sh.show_grid = ws.sheet_view.showGridLines is not False
            if ws.sheet_view.zoomScale:
                sh.zoom = max(0.25, min(4.0, ws.sheet_view.zoomScale / 100))
        except AttributeError:
            pass
        sh.recompute_extent()

    try:
        wb.active = max(0, min(book.index(book.active), len(wb.sheets) - 1)) if book.active in sheets else 0
    except (ValueError, AttributeError):
        wb.active = 0
    if not wb.sheets:
        wb.sheets.append(Sheet(wb, "Sheet1"))
    if controls_ws is not None:
        _read_controls(controls_ws, wb)

    if needs_cache:
        _load_cached_values(path, wb, needs_cache)
    wb.rebuild_dependencies()
    wb.recalc(full=True)
    return wb


def _read_freeze(ws, sh):
    """Excel stores frozen panes as a split (xSplit/ySplit) measured from the view's topLeftCell."""
    try:
        pane = ws.sheet_view.pane
    except AttributeError:
        return
    if pane is None or pane.state not in ("frozen", "frozenSplit"):
        return
    tl = parse_range(ws.sheet_view.topLeftCell or "A1") or (0, 0, 0, 0)
    ys, xs = int(pane.ySplit or 0), int(pane.xSplit or 0)
    sh.freeze = (tl[0] + ys if ys else 0, tl[1] + xs if xs else 0)
    sh.freeze_origin = (tl[0] if ys else 0, tl[1] if xs else 0)


def _write_freeze(ws, sh):
    fr, fc = sh.freeze
    if not (fr or fc):
        ws.freeze_panes = None
        return
    orr, orc = sh.freeze_origin
    ys, xs = (fr - orr if fr else 0), (fc - orc if fc else 0)
    ws.freeze_panes = f"{col_name(xs)}{ys + 1}"  # sets up the pane/selection objects
    pane = ws.sheet_view.pane
    pane.topLeftCell = f"{col_name(fc)}{fr + 1}"
    ws.sheet_view.topLeftCell = f"{col_name(orc)}{orr + 1}" if (orr or orc) else None


def _write_validation_and_cf(ws, sh):
    from openpyxl.worksheet.cell_range import MultiCellRange
    if sh.xl_dv or getattr(ws, "data_validations", None) is not None:
        try:
            dvs = []
            for dv, rects in sh.xl_dv:
                dv.sqref = MultiCellRange(" ".join(range_addr(*r) for r in rects))
                dvs.append(dv)
            ws.data_validations.dataValidation = dvs
        except AttributeError:
            pass
    if getattr(sh, "cf_complete", False):
        from openpyxl.formatting.formatting import ConditionalFormattingList
        cfl = ConditionalFormattingList()
        for rule in sh.cond_formats:
            if rule.xl_rule is not None and rule.rects:
                cfl.add(" ".join(range_addr(*r) for r in rule.rects), rule.xl_rule)
        ws.conditional_formatting = cfl


CONTROLS_SHEET = "_EkxelControls"


def _write_controls(book, wb):
    """Form controls (sliders/spinners) go in a very hidden sheet: one JSON object per row.
    Excel doesn't show it; Ekxel reads it back. Removed when there are no controls."""
    import json
    if CONTROLS_SHEET in book.sheetnames:
        book.remove(book[CONTROLS_SHEET])
    rows = []
    for sh in wb.sheets:
        for c in getattr(sh, "controls", []):
            rows.append(json.dumps({"sheet": sh.name, "kind": c["kind"], "place": list(c["place"]),
                                    "link": list(c["link"]), "min": c["min"], "max": c["max"], "step": c["step"]}))
    if not rows:
        return
    ws = book.create_sheet(CONTROLS_SHEET)
    ws.sheet_state = "veryHidden"
    ws["A1"] = "Ekxel form controls (sliders / spin buttons), one JSON object per row. Safe to delete."
    for i, row in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=row).data_type = "s"


def _read_controls(ws, wb):
    import json
    from .ui.controls import make_control
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        try:
            d = json.loads(row[0])
            sh = wb.get_sheet(d["sheet"])
            if sh is not None:
                sh.controls = sh.controls + [make_control(d["kind"], tuple(d["place"]), tuple(d["link"]),
                                                          d["min"], d["max"], d["step"])]
        except (TypeError, ValueError, KeyError):
            continue


def _update_defined_names(book, wb):
    try:
        for name, dn in book.defined_names.items():
            new = wb.names.get(name.upper())
            if new is not None and dn.attr_text != new:
                dn.attr_text = new
    except AttributeError:
        pass


def _uses_names(f):
    def walk(n):
        if n[0] == "name":
            return True
        if n[0] == "func":
            return any(walk(a) for a in n[2])
        if n[0] == "bin":
            return walk(n[2]) or walk(n[3])
        if n[0] in ("neg", "pct"):
            return walk(n[1])
        return False
    return walk(f.ast)


def _load_cached_values(path, wb, wanted):
    """Use Excel's cached results for formulas we can't evaluate ourselves."""
    import openpyxl
    try:
        vbook = openpyxl.load_workbook(_source(path), data_only=True, read_only=True)
    except Exception:
        return
    by_sheet = {}
    for sh, k in wanted:
        by_sheet.setdefault(sh.name, []).append(k)
    try:
        for name, keys in by_sheet.items():
            if name not in vbook.sheetnames:
                continue
            ws = vbook[name]
            sh = wb.get_sheet(name)
            for k in keys:
                r, c = k >> 14, k & 0x3FFF
                try:
                    v = ws.cell(row=r + 1, column=c + 1).value
                except Exception:
                    continue
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    v = float(v)
                elif isinstance(v, (dt.datetime, dt.date, dt.time)):
                    v = datetime_to_serial(v)
                elif isinstance(v, str) and errors.from_code(v):
                    v = errors.from_code(v)
                f = sh.formulas.get(k)
                if f is not None and v is not None:
                    f.fallback = v
                    f.unknown = True
    finally:
        vbook.close()


# ---------------------------------------------------------------- saving

def _style_array(style, book, cache, scratch, default_font=("Calibri", 11.0)):
    """Style -> openpyxl StyleArray, building it once on a scratch cell."""
    sa = cache.get(style)
    if sa is not None:
        return sa
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.styles.cell_style import StyleArray
    cell = scratch
    cell._style = StyleArray()
    font_kw = dict(name=style.font or default_font[0], sz=style.size or default_font[1], b=style.bold,
                   i=style.italic, u="single" if style.underline else None, strike=style.strike)
    if style.color:
        font_kw["color"] = "FF" + style.color[1:]
    cell.font = Font(**font_kw)
    if style.fill:
        cell.fill = PatternFill(fill_type="solid", fgColor="FF" + style.fill[1:], bgColor="FF" + style.fill[1:])
    else:
        cell.fill = PatternFill()
    cell.alignment = Alignment(horizontal=style.halign, vertical=style.valign,
                               wrap_text=style.wrap or None, indent=style.indent or 0)
    sides = []
    for b in style.border:
        if b:
            sides.append(Side(style=b[0], color="FF" + (b[1] or "#000000")[1:]))
        else:
            sides.append(Side())
    cell.border = Border(left=sides[0], top=sides[1], right=sides[2], bottom=sides[3])
    cell.number_format = style.numfmt or "General"
    sa = copy(cell._style)
    cache[style] = sa
    return sa


def save_xlsx(wb, path):
    import openpyxl
    from openpyxl.worksheet.cell_range import MultiCellRange
    from openpyxl.worksheet.dimensions import DimensionHolder

    reuse = wb.xl_book is not None
    ext = os.path.splitext(path)[1].lower()
    default_font = getattr(wb, "xl_default_font", ("Calibri", 11.0)) if reuse else ("Calibri", 11.0)
    if reuse:
        book = wb.xl_book
        style_cache = {}
    else:
        book = openpyxl.Workbook()
        style_cache = {}
        for ws in list(book.worksheets):
            book.remove(ws)
    if ext == ".xlsx" and getattr(book, "vba_archive", None) is not None:
        book.vba_archive = None
        book.is_template = False

    # sheets: reuse, create, remove, order
    wanted = []
    for i, sh in enumerate(wb.sheets):
        ws = sh.xl if (sh.xl is not None and sh.xl in book.worksheets) else None
        if ws is None:
            ws = book.create_sheet(f"__new{i}")
            sh.xl = ws
        wanted.append(ws)
    for ws in list(book.worksheets):
        if ws not in wanted:
            book.remove(ws)
    for i, ws in enumerate(wanted):
        ws.title = f"__tmp{i}"
    for sh, ws in zip(wb.sheets, wanted):
        ws.title = sh.name
    others = [s for s in book._sheets if s not in wanted]
    book._sheets = wanted + others

    scratch_ws = wanted[0]
    scratch = None
    for sh, ws in zip(wb.sheets, wanted):
        ws.merged_cells = MultiCellRange()
        old_cells = ws._cells
        ws._cells = {}
        for oc in old_cells.values():  # release comments so they can bind to new cells
            if getattr(oc, "comment", None) is not None:
                oc.comment = None
        # notes and hyperlinks, at their current (possibly moved) positions
        for k, (com, link) in sh.notes.items():
            nc = ws.cell(row=(k >> 14) + 1, column=(k & 0x3FFF) + 1)
            if com is not None:
                if getattr(com, "parent", None) is not None:
                    com = copy(com)
                nc.comment = com
            if link is not None:
                link.ref = nc.coordinate
                nc._hyperlink = link
        if scratch is None:
            scratch = openpyxl.cell.cell.Cell(scratch_ws, row=1, column=1)
        has_styles = sh.styles
        xl_styles = sh.xl_styles if reuse else {}
        keys = set(sh.values) | set(sh.formulas) | set(has_styles) | set(xl_styles)
        for k in sorted(keys):
            r, c = k >> 14, k & 0x3FFF
            cell = ws.cell(row=r + 1, column=c + 1)
            f = sh.formulas.get(k)
            if f is not None:
                cell.value = to_file_formula(f.text)
            elif k in sh.values:
                v = sh.values[k]
                if isinstance(v, XLError):
                    cell.value = v.code
                    cell.data_type = "e"
                elif isinstance(v, bool):
                    cell.value = v
                elif isinstance(v, float):
                    cell.value = int(v) if v.is_integer() and abs(v) < 2 ** 53 else v
                else:
                    cell.value = v
                    if isinstance(v, str):
                        cell.data_type = "s"
            st = has_styles.get(k)
            orig = xl_styles.get(k)
            if orig is not None and orig[0] == (st or DEFAULT_STYLE):
                cell._style = copy(orig[1])
            elif st is not None:
                cell._style = copy(_style_array(st, book, style_cache, scratch, default_font))
        # geometry
        ws.column_dimensions = DimensionHolder(worksheet=ws, default_factory=ws._add_column)
        ws.row_dimensions = DimensionHolder(worksheet=ws, default_factory=ws._add_row)
        for c in sorted(set(sh.col_widths) | sh.hidden_cols):
            cd = ws.column_dimensions[col_name(c)]
            px = sh.col_widths.get(c, DEFAULT_COL_WIDTH)
            cd.width = round(max(0, (px - 5) / 7), 2)
            if c in sh.hidden_cols:
                cd.hidden = True
        for r in sorted(set(sh.row_heights) | sh.hidden_rows | sh.filter_hidden):
            rd = ws.row_dimensions[r + 1]
            if r in sh.row_heights:
                rd.ht = round(sh.row_heights[r] * 72 / 96, 2)
            if r in sh.hidden_rows or r in sh.filter_hidden:
                rd.hidden = True
        for m in sh.merges:
            ws.merge_cells(start_row=m[0] + 1, start_column=m[1] + 1, end_row=m[2] + 1, end_column=m[3] + 1)
        _write_freeze(ws, sh)
        _write_validation_and_cf(ws, sh)
        ws.auto_filter.ref = range_addr(*sh.autofilter) if sh.autofilter else None
        try:
            ws.sheet_view.showGridLines = bool(sh.show_grid)
            ws.sheet_view.zoomScale = int(round(sh.zoom * 100)) if sh.zoom != 1.0 else None
            ws.sheet_properties.tabColor = ("FF" + sh.tab_color[1:]) if sh.tab_color else None
        except (AttributeError, ValueError):
            pass
    try:
        book.active = wb.active
        for i, ws in enumerate(wanted):
            ws.sheet_view.tabSelected = (i == wb.active)
    except (ValueError, IndexError, AttributeError):
        pass
    _write_controls(book, wb)
    if reuse:
        _update_defined_names(book, wb)
    tmp = path + ".tmp~"
    book.save(tmp)
    os.replace(tmp, path)
    wb.xl_book = book
