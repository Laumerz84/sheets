"""Charts in .xlsx files.

Saving: each Ekxel chart is written as a real Excel chart (openpyxl.chart) next to the cells, so Excel
shows it with the same type, data links, title, axes, legend, colours and position. Ekxel's own
complete setup of every chart is kept as JSON in the hidden `_EkxelControls` sheet (see io_xlsx),
which is what Ekxel reads back when the file is reopened.

Charts that came from other programs are not shown or edited by Ekxel; openpyxl carries them through
a save with some loss of formatting (`foreign_charts` counts them for the warning)."""
import re
import zipfile
from xml.etree.ElementTree import Element

from . import charts as C

EMU = 9525          # EMU per pixel at 96 dpi
_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def _ref(text):
    """'=Sheet1!$A$1:$A$9' -> the same without the '=' (rows clipped to what an xlsx file can hold)."""
    p = C.parse_ref(text)
    if p is None:
        return text[1:] if text.startswith("=") else text
    r1, c1, r2, c2 = p[1]
    return C.make_ref(p[0], (min(r1, 1048575), c1, min(r2, 1048575), c2))[1:]


def _xref(text):
    """An openpyxl Reference for '=Sheet!$A$1:$A$9' (sheet names with quotes or spaces come out right)."""
    from openpyxl.chart import Reference
    name, rect = C.parse_ref(text)
    r1, c1, r2, c2 = rect
    ref = Reference(range_string="'x'!" + C.range_addr(min(r1, 1048575), c1, min(r2, 1048575), c2))
    ref.worksheet.title = name
    return ref


def _hex(c):
    return c.lstrip("#").upper()


def _resolved_title(ch, names):
    t = ch.get("title")
    if t is None:
        return None
    if t == "":
        return names[0] if (len(names) == 1 or C.family(ch["type"]) == "pie") and names[0] else "Chart Title"
    return t


def _series_names(wb, ch, home):
    out = []
    for i, s in enumerate(ch["series"]):
        n = s.get("name") or ""
        if n.startswith("="):
            got = C.ref_rect(wb, n, home.name)
            v = got[0].value(got[1][0], got[1][1]) if got else None
            n = "" if v is None else str(v)
        out.append(n or f"Series{i + 1}")
    return out


def _font(sz_pt, color, bold=False):
    from openpyxl.drawing.text import CharacterProperties
    return CharacterProperties(sz=int(sz_pt * 100), b=bold, solidFill=_hex(color), latin=_latin())


def _latin():
    from openpyxl.drawing.text import Font
    return Font(typeface="Calibri")


def _txpr(sz_pt, color, bold=False):
    """Text properties (size, colour) for an axis / legend / data labels."""
    from openpyxl.chart.text import RichText
    from openpyxl.drawing.text import Paragraph, ParagraphProperties
    cp = _font(sz_pt, color, bold)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _title(text, sz_pt, color):
    """A chart / axis title in Ekxel's look (grey Calibri)."""
    from openpyxl.chart.text import RichText, Text
    from openpyxl.chart.title import Title
    from openpyxl.drawing.text import Paragraph, ParagraphProperties, RegularTextRun
    cp = _font(sz_pt, color)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def _gridlines(color):
    from openpyxl.chart.axis import ChartLines
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    return ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=_hex(color), w=9525)))


def _axis_line(color):
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    return GraphicalProperties(ln=LineProperties(solidFill=_hex(color), w=9525))


def _no_auto_title(chart):
    """Excel shows a one-series chart's name as its title unless autoTitleDeleted is set."""
    orig = chart._write

    def write():
        tree = orig()
        c = tree.find(f"{{{_NS}}}chart")
        pa = c.find(f"{{{_NS}}}plotArea") if c is not None else None
        if c is not None and pa is not None and c.find(f"{{{_NS}}}autoTitleDeleted") is None:
            el = Element(f"{{{_NS}}}autoTitleDeleted")
            el.set("val", "1")
            c.insert(list(c).index(pa), el)
        return tree
    chart._write = write


def _make_series(ch, s, i, name, color, kind, scatter=False):
    from openpyxl.chart import Series
    from openpyxl.chart.data_source import AxDataSource, StrRef
    from openpyxl.chart.series import SeriesLabel
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    vals = _xref(s["values"])
    cats = _ref(s["cats"]) if s.get("cats") else None
    if scatter:
        ser = Series(vals, xvalues=_xref(s["cats"])) if cats else Series(vals)
    else:
        ser = Series(vals)
        if cats:
            ser.cat = AxDataSource(strRef=StrRef(f=cats))
    nm = s.get("name") or ""
    if nm.startswith("="):
        ser.tx = SeriesLabel(strRef=StrRef(f=_ref(nm)))
    else:
        ser.tx = SeriesLabel(v=name)
    col = _hex(color)
    if kind in ("line", "scatter_line", "scatter"):
        ser.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=col, w=28575))
    else:
        ser.graphicalProperties = GraphicalProperties(solidFill=col)
    return ser


def _marker(ser, color, on):
    from openpyxl.chart.marker import Marker
    from openpyxl.chart.shapes import GraphicalProperties
    if on:
        ser.marker = Marker(symbol="circle", size=7, spPr=GraphicalProperties(solidFill=_hex(color)))
        ser.marker.spPr.line.solidFill = _hex(color)
    else:
        ser.marker = Marker(symbol="none")


def _build_one(ch, kind, idxs, colors, names, grouping, secondary, npts=12):
    """An openpyxl chart object of `kind` for the series with indexes `idxs`."""
    from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, ScatterChart
    series = ch["series"]
    if kind in ("col", "bar"):
        c = BarChart()
        c.type = "col" if kind == "col" else "bar"
        c.grouping = {"clustered": "clustered", "stacked": "stacked", "percent": "percentStacked"}[grouping]
        if grouping != "clustered":
            c.overlap = 100
        c.gapWidth = int(ch.get("gap") or (150 if grouping != "clustered" else 219))
        for i in idxs:
            c.series.append(_make_series(ch, series[i], i, names[i], colors[i], "col"))
    elif kind in ("line", "line_markers"):
        c = LineChart()
        c.grouping = "standard"
        for i in idxs:
            ser = _make_series(ch, series[i], i, names[i], colors[i], "line")
            mk = series[i].get("marker")
            _marker(ser, colors[i], (kind == "line_markers") if mk is None else bool(mk))
            ser.smooth = bool(series[i].get("smooth"))
            c.series.append(ser)
    elif kind == "area":
        c = AreaChart()
        c.grouping = {"clustered": "standard", "stacked": "stacked", "percent": "percentStacked"}[grouping]
        for i in idxs:
            c.series.append(_make_series(ch, series[i], i, names[i], colors[i], "area"))
    elif kind in ("scatter", "scatter_lines", "scatter_smooth"):
        c = ScatterChart()
        c.scatterStyle = "lineMarker"
        for i in idxs:
            ser = _make_series(ch, series[i], i, names[i], colors[i], "scatter", scatter=True)
            _marker(ser, colors[i], True)
            if kind == "scatter":
                ser.graphicalProperties.line.noFill = True
            ser.smooth = kind == "scatter_smooth" or bool(series[i].get("smooth"))
            c.series.append(ser)
    elif kind in ("pie", "doughnut"):
        c = PieChart() if kind == "pie" else DoughnutChart()
        if kind == "doughnut":
            c.holeSize = int(ch.get("hole", 60))
        from openpyxl.chart.series import DataPoint
        from openpyxl.chart.shapes import GraphicalProperties
        for i in idxs:
            ser = _make_series(ch, series[i], i, names[i], colors[i], "area")
            n = npts
            pal = C.palette_colors(ch.get("style", 1), max(n, 1))
            ser.dPt = [DataPoint(idx=j, spPr=GraphicalProperties(solidFill=_hex(pal[j]))) for j in range(n)]
            c.series.append(ser)
            if kind == "pie":
                break
    else:
        raise ValueError(kind)
    if secondary and kind not in ("pie", "doughnut"):
        c.y_axis.axId = 200
        c.y_axis.crosses = "max"
        c.y_axis.majorGridlines = None
    return c


def to_openpyxl(wb, home, ch):
    """The openpyxl chart object for Ekxel chart `ch` (which lives on sheet `home`)."""
    from openpyxl.chart.label import DataLabelList
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
    from openpyxl.chart.data_source import NumFmt
    from openpyxl.drawing.line import LineProperties

    fam = C.family(ch["type"])
    # series whose data sheet is gone (deleted) can't be written
    ch = dict(ch, series=[s for s in ch["series"] if C.ref_rect(wb, s.get("values", ""), home.name)])
    if not ch["series"]:
        raise ValueError("no data")
    names = _series_names(wb, ch, home)
    series = ch["series"]
    pie = fam in ("pie", "doughnut")
    colors = C.palette_colors(ch.get("style", 1), len(series))
    colors = [s.get("color") or colors[i] for i, s in enumerate(series)]
    grouping = C.grouping(ch["type"]) if fam in ("col", "bar", "area") else "clustered"
    npts = 12
    if pie:     # slice colours need the slice count
        npts = max([len(s["y"]) for s in C.resolve(wb, ch, home)["series"]] or [12])
    if fam == "combo":
        groups = {}
        for i, s in enumerate(series):
            kind = s.get("type") or "col"
            if kind.startswith("scatter"):
                kind = "line"
            groups.setdefault((kind, bool(s.get("secondary"))), []).append(i)
        order = sorted(groups, key=lambda k: (k[1], 0 if k[0] in ("col", "area") else 1))
        charts = [_build_one(ch, k, groups[(k, sec)], colors, names, "clustered", sec) for k, sec in order]
        chart = charts[0]
        for extra in charts[1:]:
            chart += extra
        sec_chart = next((c for c, (k, s) in zip(charts, order) if s), None)
    else:
        kind = ch["type"] if fam in ("scatter", "line") else fam
        chart = _build_one(ch, kind, list(range(len(series))), colors, names, grouping, False, npts)
        sec_chart = None
    # title / legend / chart area
    title = _resolved_title(ch, names)
    dark = C.is_dark(ch.get("fill") or "#FFFFFF")
    txt = "#D9D9D9" if dark else "#595959"
    grid = "#595959" if dark else "#D9D9D9"
    chart.title = _title(title, 14, txt) if title is not None else None
    if title is None:
        _no_auto_title(chart)
    chart.style = None
    if ch["legend"] == "none":
        chart.legend = None
    else:
        chart.legend.position = {"right": "r", "left": "l", "top": "t", "bottom": "b"}[ch["legend"]]
        chart.legend.txPr = _txpr(9, txt)
        chart.legend.overlay = False
    if ch.get("fill") or not ch.get("border", True):
        gp = GraphicalProperties(solidFill=_hex(ch.get("fill") or "#FFFFFF"))
        if not ch.get("border", True):
            gp.line.noFill = True
        else:
            gp.line.solidFill = "D9D9D9"
        chart.graphical_properties = gp
    chart.roundedCorners = False
    if not pie:
        xa, ya = chart.x_axis, chart.y_axis
        xa.delete = False
        ya.delete = False
        if ch["x_title"]:
            xa.title = _title(ch["x_title"], 10, txt)
        if ch["y_title"]:
            ya.title = _title(ch["y_title"], 10, txt)
        for ax in (xa, ya):
            ax.txPr = _txpr(9, txt)
        xa.spPr = _axis_line(grid)
        ya.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
        ya.majorGridlines = _gridlines(grid) if ch["grid_y"] else None
        xa.majorGridlines = _gridlines(grid) if ch["grid_x"] else None
        if ch.get("x_reverse"):
            xa.scaling.orientation = "maxMin"
        for attr, val in (("min", ch["y_min"]), ("max", ch["y_max"])):
            if val is not None:
                setattr(ya.scaling, attr, val)
        if ch["y_major"]:
            ya.majorUnit = ch["y_major"]
        if ch.get("log_y"):
            ya.scaling.logBase = 10
        if ch.get("y_fmt"):
            ya.numFmt = NumFmt(formatCode=ch["y_fmt"], sourceLinked=False)
        elif grouping == "percent":
            ya.numFmt = NumFmt(formatCode="0%", sourceLinked=False)
        else:
            ya.numFmt = NumFmt(formatCode="General", sourceLinked=True)   # like Excel's "Linked to source"
        if fam == "scatter":
            for attr, val in (("min", ch["x_min"]), ("max", ch["x_max"])):
                if val is not None:
                    setattr(xa.scaling, attr, val)
        if sec_chart is not None:
            sy = sec_chart.y_axis
            sy.delete = False
            sy.txPr = _txpr(9, txt)
            sy.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
            sy.numFmt = NumFmt(formatCode="General", sourceLinked=True)
            if ch["y2_title"]:
                sy.title = _title(ch["y2_title"], 10, txt)
            for attr, val in (("min", ch["y2_min"]), ("max", ch["y2_max"])):
                if val is not None:
                    setattr(sy.scaling, attr, val)
    if ch["labels"] or any(s.get("labels") for s in series):
        dl = DataLabelList()
        dl.showVal = not (pie and ch.get("label_pct"))
        dl.showPercent = bool(pie and ch.get("label_pct"))
        dl.showSerName = dl.showCatName = dl.showLegendKey = False
        dl.txPr = _txpr(9, txt)
        chart.dataLabels = dl
    # where it sits (two-cell anchor like Excel's "move and size with cells")
    f, t = ch["from"], ch["to"]
    a = TwoCellAnchor(editAs="twoCell")
    a._from = AnchorMarker(col=f[1], row=f[0], colOff=f[2] * EMU, rowOff=f[3] * EMU)
    a.to = AnchorMarker(col=t[1], row=t[0], colOff=t[2] * EMU, rowOff=t[3] * EMU)
    chart.anchor = a
    chart._ekxel = True
    return chart


# ================================================================ saving / loading hooks
def write_charts(wb):
    """Put Ekxel's charts on the openpyxl sheets (replacing the ones written by the last save)."""
    for sh in wb.sheets:
        ws = sh.xl
        if ws is None:
            continue
        ws._charts = [c for c in ws._charts if not getattr(c, "_ekxel", False)]
        for ch in getattr(sh, "charts", []):
            try:
                ws.add_chart(to_openpyxl(wb, sh, ch))
            except ValueError:      # nothing to plot (its data is gone): leave the chart out of the Excel copy
                pass
            except Exception:       # one odd chart must not stop the whole save
                import traceback
                traceback.print_exc()


def json_rows(wb):
    """(sheet name, chart json) for every chart, for the hidden sheet."""
    out = []
    for sh in wb.sheets:
        for ch in getattr(sh, "charts", []):
            out.append((sh.name, C.to_json(ch)))
    return out


def read_json(wb, sheet_name, d):
    sh = wb.get_sheet(sheet_name)
    if sh is None:
        return
    sh.charts = sh.charts + [C.from_json(d)]


def drop_written_copies(wb):
    """After loading: remove the openpyxl copies of charts Ekxel restored from its hidden sheet (they are
    the last charts of the sheet, in order, since Ekxel writes its own after any foreign ones)."""
    for sh in wb.sheets:
        n = len(getattr(sh, "charts", []))
        ws = sh.xl
        if n and ws is not None and len(ws._charts) >= n:
            ws._charts = ws._charts[:len(ws._charts) - n]


def foreign_charts(path, wb):
    """How many charts in the file Ekxel didn't make (and so can't show)."""
    try:
        with zipfile.ZipFile(path) as z:
            n = sum(1 for name in z.namelist() if re.fullmatch(r"xl/charts/chart\d+\.xml", name))
    except (zipfile.BadZipFile, OSError):
        return 0
    return max(0, n - sum(len(getattr(s, "charts", [])) for s in wb.sheets))
