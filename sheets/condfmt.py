"""Conditional formatting: rules read from xlsx files or made in Ekxel (Home > Conditional
Formatting), how they're evaluated for display, and how Ekxel-made rules are written back.
Rules are never changed in place (undo snapshots keep references to them): edits build new ones."""
import colorsys

from .errors import XLError
from .formula import ParseError, parse, shift_formula
from .values import compare, is_num, to_bool


class CFRule:
    __slots__ = ("rects", "kind", "op", "formulas", "fill", "color", "bold", "italic",
                 "scale", "bar", "text", "rank", "percent", "bottom", "above", "stop", "priority", "_asts",
                 "xl_rule", "made")

    def __init__(self, rects, kind, priority=0):
        self.rects = rects
        self.kind = kind
        self.op = None
        self.formulas = []
        self.fill = self.color = None
        self.bold = self.italic = None
        self.scale = None      # [(cfvo_type, value, color)]
        self.bar = None        # (color, cfvo_min, cfvo_max)
        self.text = None
        self.rank = 10
        self.percent = False
        self.bottom = False
        self.above = True
        self.stop = False
        self.priority = priority
        self._asts = {}
        self.xl_rule = None     # the openpyxl Rule, written back on save
        self.made = False       # made (or edited) in Ekxel rather than read from the file

    def contains(self, r, c):
        for r1, c1, r2, c2 in self.rects:
            if r1 <= r <= r2 and c1 <= c <= c2:
                return True
        return False

    def anchor(self):
        return min(x[0] for x in self.rects), min(x[1] for x in self.rects)


def _interp(c1, c2, t):
    a = [int(c1[i:i + 2], 16) for i in (1, 3, 5)]
    b = [int(c2[i:i + 2], 16) for i in (1, 3, 5)]
    return "#%02X%02X%02X" % tuple(round(x + (y - x) * t) for x, y in zip(a, b))


class CFEngine:
    """Evaluates rules for cells with per-paint caching."""

    def __init__(self, sheet):
        self.sheet = sheet
        self.cache = {}
        self.range_stats = {}

    def clear(self):
        self.cache = {}
        self.range_stats = {}

    def _eval(self, rule, idx, r, c):
        """Evaluate rule.formulas[idx] relative to the rule's anchor for cell (r, c)."""
        ar, ac = rule.anchor()
        dr, dc = r - ar, c - ac
        key = (idx, dr, dc)
        ast = rule._asts.get(key)
        if ast is None:
            text = rule.formulas[idx]
            if not text.startswith("="):
                text = "=" + text
            try:
                ast = parse(shift_formula(text, dr, dc))
            except (ParseError, Exception):
                ast = ("err", None)
            if len(rule._asts) < 20000:
                rule._asts[key] = ast
        if ast[0] == "err" and ast[1] is None:
            return None
        return self.sheet.wb.evaluator.run(ast, (self.sheet, r, c))

    def _numbers(self, rule):
        st = self.range_stats.get(id(rule))
        if st is None:
            sh = self.sheet
            nums = []
            texts = {}
            for r1, c1, r2, c2 in rule.rects:
                r2 = min(r2, sh.max_row)
                c2 = min(c2, sh.max_col)
                for r in range(r1, r2 + 1):
                    for c in range(c1, c2 + 1):
                        v = sh.value(r, c)
                        if is_num(v):
                            nums.append(v)
                        if v is not None and v != "":
                            t = str(v).lower()
                            texts[t] = texts.get(t, 0) + 1
            nums.sort()
            st = (nums, texts)
            self.range_stats[id(rule)] = st
        return st

    def _cfvo(self, rule, kind, val, nums, default):
        if not nums:
            return default
        try:
            if kind == "min":
                return nums[0]
            if kind == "max":
                return nums[-1]
            if kind == "num":
                return float(val)
            if kind == "percent":
                lo, hi = nums[0], nums[-1]
                return lo + (hi - lo) * float(val) / 100
            if kind == "percentile":
                p = float(val) / 100
                pos = p * (len(nums) - 1)
                i = int(pos)
                j = min(i + 1, len(nums) - 1)
                return nums[i] + (nums[j] - nums[i]) * (pos - i)
        except (TypeError, ValueError):
            pass
        return default

    def style_for(self, rules, r, c):
        """-> dict of overrides (fill, color, bold, italic, bar) or None."""
        k = (r, c)
        if k in self.cache:
            return self.cache[k]
        out = None
        v = None
        got_v = False
        for rule in rules:
            if not rule.contains(r, c):
                continue
            if not got_v:
                v = self.sheet.value(r, c)
                got_v = True
            try:
                hit = self._match(rule, v, r, c)
            except XLError:
                hit = False
            if not hit:
                continue
            if out is None:
                out = {}
            if rule.kind == "colorScale":
                out.setdefault("fill", hit)
            elif rule.kind == "dataBar":
                out.setdefault("bar", hit)
            else:
                for f in ("fill", "color", "bold", "italic"):
                    val = getattr(rule, f)
                    if val is not None:
                        out.setdefault(f, val)
            if rule.stop:
                break
        if len(self.cache) > 50000:
            self.cache = {}
        self.cache[k] = out
        return out

    def _match(self, rule, v, r, c):
        kind = rule.kind
        if kind == "cellIs":
            if v is None:
                v = 0.0
            if isinstance(v, XLError):
                return False
            vals = [self._eval(rule, i, r, c) for i in range(len(rule.formulas))]
            if not vals or any(isinstance(x, XLError) or x is None for x in vals):
                return False
            op = rule.op
            if op in ("between", "notBetween") and len(vals) >= 2:
                lo, hi = sorted(vals[:2], key=lambda x: (0, x) if is_num(x) else (1, str(x)))
                inside = compare(v, lo) >= 0 and compare(v, hi) <= 0
                return inside if op == "between" else not inside
            x = compare(v, vals[0])
            return {"equal": x == 0, "notEqual": x != 0, "greaterThan": x > 0, "lessThan": x < 0,
                    "greaterThanOrEqual": x >= 0, "lessThanOrEqual": x <= 0}.get(op, False)
        if kind == "expression":
            res = self._eval(rule, 0, r, c)
            if res is None or isinstance(res, XLError):
                return False
            try:
                return to_bool(res)
            except XLError:
                return False
        if kind in ("containsText", "notContainsText", "beginsWith", "endsWith"):
            if v is None or rule.text is None:
                return kind == "notContainsText"
            t, n = str(v).lower(), rule.text.lower()
            return {"containsText": n in t, "notContainsText": n not in t,
                    "beginsWith": t.startswith(n), "endsWith": t.endswith(n)}[kind]
        if kind in ("containsBlanks", "notContainsBlanks"):
            blank = v is None or (isinstance(v, str) and not v.strip())
            return blank if kind == "containsBlanks" else not blank
        if kind in ("containsErrors", "notContainsErrors"):
            return isinstance(v, XLError) == (kind == "containsErrors")
        if kind == "colorScale" and is_num(v) and rule.scale:
            nums, _ = self._numbers(rule)
            pts = [(self._cfvo(rule, t, val, nums, 0.0), col) for t, val, col in rule.scale]
            if len(pts) < 2:
                return False
            if v <= pts[0][0]:
                return pts[0][1]
            for (a, ca), (b, cb) in zip(pts, pts[1:]):
                if v <= b:
                    return _interp(ca, cb, 0 if b == a else (v - a) / (b - a))
            return pts[-1][1]
        if kind == "dataBar" and is_num(v) and rule.bar:
            nums, _ = self._numbers(rule)
            color, lo_t, hi_t = rule.bar
            lo = self._cfvo(rule, lo_t[0], lo_t[1], nums, 0.0)
            hi = self._cfvo(rule, hi_t[0], hi_t[1], nums, 1.0)
            if hi == lo:
                return (1.0, color)
            return (max(0.0, min(1.0, (v - lo) / (hi - lo))), color)
        if kind == "top10" and is_num(v):
            nums, _ = self._numbers(rule)
            n = rule.rank
            if rule.percent:
                n = max(1, int(len(nums) * rule.rank / 100))
            if not nums:
                return False
            if rule.bottom:
                return v <= nums[min(n, len(nums)) - 1]
            return v >= nums[-min(n, len(nums))]
        if kind == "aboveAverage" and is_num(v):
            nums, _ = self._numbers(rule)
            if not nums:
                return False
            avg = sum(nums) / len(nums)
            return v > avg if rule.above else v < avg
        if kind in ("duplicateValues", "uniqueValues") and v not in (None, ""):
            _, texts = self._numbers(rule)
            cnt = texts.get(str(v).lower(), 0)
            return cnt > 1 if kind == "duplicateValues" else cnt == 1
        return False


def load_rules(ws, color_fn):
    """openpyxl worksheet -> ([CFRule] sorted by priority, complete).
    `complete` is False when some rule couldn't be read; the file's own
    formatting is then left untouched on save."""
    from .refs import parse_range
    out = []
    complete = True
    try:
        entries = list(ws.conditional_formatting)
    except Exception:
        return out, False
    for cf in entries:
        rects = []
        for rng in str(cf.sqref).split():
            b = parse_range(rng)
            if b:
                rects.append(b)
            else:
                complete = False
        if not rects:
            complete = False
            continue
        for rule in cf.rules:
            kind = rule.type
            cr = CFRule(list(rects), kind, rule.priority or 0)
            cr.xl_rule = rule
            cr.op = rule.operator
            cr.formulas = [str(f) for f in (rule.formula or [])]
            cr.text = rule.text
            cr.stop = bool(rule.stopIfTrue)
            if rule.rank is not None:
                cr.rank = int(rule.rank)
            cr.percent = bool(rule.percent)
            cr.bottom = bool(rule.bottom)
            cr.above = rule.aboveAverage is not False
            dxf = rule.dxf
            if dxf is not None:
                if dxf.fill is not None:
                    fill = dxf.fill
                    col = color_fn(getattr(fill, "bgColor", None)) or color_fn(getattr(fill, "fgColor", None))
                    if getattr(fill, "fill_type", None) not in (None, "none") or col:
                        cr.fill = col
                if dxf.font is not None:
                    cr.color = color_fn(dxf.font.color)
                    if dxf.font.b is not None:
                        cr.bold = bool(dxf.font.b)
                    if dxf.font.i is not None:
                        cr.italic = bool(dxf.font.i)
            if kind == "colorScale" and rule.colorScale is not None:
                cs = rule.colorScale
                cols = [color_fn(c) or "#FFFFFF" for c in cs.color]
                cr.scale = [(v.type, v.val, col) for v, col in zip(cs.cfvo, cols)]
            if kind == "dataBar" and rule.dataBar is not None:
                db = rule.dataBar
                cfvo = list(db.cfvo) + [None, None]
                lo = (cfvo[0].type, cfvo[0].val) if cfvo[0] is not None else ("min", None)
                hi = (cfvo[1].type, cfvo[1].val) if cfvo[1] is not None else ("max", None)
                cr.bar = (color_fn(db.color) or "#638EC6", lo, hi)
            if kind in ("containsText", "notContainsText", "beginsWith", "endsWith") and not cr.text and cr.formulas:
                cr.text = None
            out.append(cr)
    out.sort(key=lambda x: x.priority)
    return out, complete


# ---------------------------------------------------------------- making rules (Ekxel's own)
# Excel's preset looks for "Highlight Cells Rules": (label, fill, font colour)
PRESETS = [("Light Red Fill with Dark Red Text", "#FFC7CE", "#9C0006"),
           ("Yellow Fill with Dark Yellow Text", "#FFEB9C", "#9C5700"),
           ("Green Fill with Dark Green Text", "#C6EFCE", "#006100"),
           ("Light Red Fill", "#FFC7CE", None),
           ("Red Text", None, "#9C0006")]
COLOR_SCALES = [("Green - Yellow - Red", ["#63BE7B", "#FFEB84", "#F8696B"]),
                ("Red - Yellow - Green", ["#F8696B", "#FFEB84", "#63BE7B"]),
                ("Green - White - Red", ["#63BE7B", "#FFFFFF", "#F8696B"]),
                ("Red - White - Green", ["#F8696B", "#FFFFFF", "#63BE7B"]),
                ("Blue - White - Red", ["#5A8AC6", "#FFFFFF", "#F8696B"]),
                ("Red - White - Blue", ["#F8696B", "#FFFFFF", "#5A8AC6"]),
                ("White - Red", ["#FFFFFF", "#F8696B"]),
                ("Green - White", ["#63BE7B", "#FFFFFF"]),
                ("Green - Yellow", ["#63BE7B", "#FFEF9C"]),
                ("Yellow - Green", ["#FFEF9C", "#63BE7B"])]
DATA_BARS = [("Blue", "#638EC6"), ("Green", "#63C384"), ("Red", "#FF555A"), ("Orange", "#FFB628"),
             ("Light Blue", "#008AEF"), ("Purple", "#D6007B")]
TEXT_KINDS = ("containsText", "notContainsText", "beginsWith", "endsWith")


def new_rule(rects, kind, **fields):
    """A fresh Ekxel-made rule; `fields` are CFRule attributes (op, formulas, fill, color, ...)."""
    rule = CFRule([tuple(x) for x in rects], kind)
    for k, v in fields.items():
        setattr(rule, k, v)
    rule.made = True
    rule.xl_rule = to_xl_rule(rule)
    return rule


def edited(rule, rects=None, **fields):
    """A copy of `rule` with some fields (and/or its rectangles) changed."""
    names = ("kind", "op", "formulas", "fill", "color", "bold", "italic", "scale", "bar", "text", "rank",
             "percent", "bottom", "above", "stop")
    vals = {n: getattr(rule, n) for n in names}
    vals.update(fields)
    kind = vals.pop("kind")
    return new_rule(rects if rects is not None else rule.rects, kind, **vals)


def value_formula(text):
    """What the user typed as a rule value -> a stored formula: '10' -> '10', 'abc' -> '"abc"',
    '=B1' -> 'B1'."""
    s = str(text).strip()
    if s.startswith("="):
        return s[1:]
    try:
        float(s)
        return s
    except ValueError:
        return '"' + s.replace('"', '""') + '"'


def formula_display(f):
    """A stored rule formula back into what a dialog shows."""
    s = str(f)
    if len(s) >= 2 and s[0] == '"' and s[-1] == '"':
        return s[1:-1].replace('""', '"')
    try:
        float(s)
        return s
    except ValueError:
        return "=" + s


def _argb(c):
    return "FF" + c.lstrip("#").upper()[-6:]


def _dxf(rule):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.styles.differential import DifferentialStyle
    font = None
    if rule.color or rule.bold is not None or rule.italic is not None:
        font = Font(color=_argb(rule.color) if rule.color else None, b=rule.bold, i=rule.italic)
    fill = PatternFill(patternType="solid", bgColor=_argb(rule.fill)) if rule.fill else None
    return DifferentialStyle(font=font, fill=fill)


def to_xl_rule(rule):
    """The openpyxl Rule Excel needs for an Ekxel-made rule."""
    from openpyxl.formatting.rule import ColorScale, DataBar, FormatObject, Rule
    from openpyxl.styles.colors import Color
    from .refs import addr
    k = rule.kind
    x = Rule(type=k, stopIfTrue=True if rule.stop else None)
    ar, ac = rule.anchor()
    cell = addr(ar, ac)
    if k in ("cellIs", "expression", "duplicateValues", "uniqueValues", "top10", "aboveAverage",
             "containsBlanks", "notContainsBlanks", "containsErrors", "notContainsErrors") + TEXT_KINDS:
        x.dxf = _dxf(rule)
    if k == "cellIs":
        x.operator = rule.op
        x.formula = list(rule.formulas)
    elif k == "expression":
        x.formula = [rule.formulas[0].lstrip("=")] if rule.formulas else []
    elif k in TEXT_KINDS:
        t = (rule.text or "").replace('"', '""')
        x.operator = k
        x.text = rule.text
        x.formula = [{"containsText": f'NOT(ISERROR(SEARCH("{t}",{cell})))',
                      "notContainsText": f'ISERROR(SEARCH("{t}",{cell}))',
                      "beginsWith": f'LEFT({cell},LEN("{t}"))="{t}"',
                      "endsWith": f'RIGHT({cell},LEN("{t}"))="{t}"'}[k]]
    elif k in ("containsBlanks", "notContainsBlanks"):
        x.formula = [f"LEN(TRIM({cell}))=0" if k == "containsBlanks" else f"LEN(TRIM({cell}))>0"]
    elif k in ("containsErrors", "notContainsErrors"):
        x.formula = [f"ISERROR({cell})" if k == "containsErrors" else f"NOT(ISERROR({cell}))"]
    elif k == "top10":
        x.rank, x.percent, x.bottom = int(rule.rank), bool(rule.percent) or None, bool(rule.bottom) or None
    elif k == "aboveAverage":
        x.aboveAverage = None if rule.above else False
    elif k == "colorScale":
        x.colorScale = ColorScale(cfvo=[FormatObject(type=t, val=v) for t, v, _ in rule.scale],
                                  color=[Color(rgb=_argb(c)) for _, _, c in rule.scale])
    elif k == "dataBar":
        color, lo, hi = rule.bar
        x.dataBar = DataBar(cfvo=[FormatObject(type=lo[0], val=lo[1]), FormatObject(type=hi[0], val=hi[1])],
                            color=Color(rgb=_argb(color)))
    return x


def scale_rule(rects, colors):
    pts = [("min", None)] + ([("percentile", 50)] if len(colors) == 3 else []) + [("max", None)]
    return new_rule(rects, "colorScale", scale=[(t, v, c) for (t, v), c in zip(pts, colors)])


def bar_rule(rects, color):
    return new_rule(rects, "dataBar", bar=(color, ("min", None), ("max", None)))


def describe(rule):
    """One line for the Manage Rules list."""
    k = rule.kind
    ops = {"between": "between", "notBetween": "not between", "equal": "=", "notEqual": "<>",
           "greaterThan": ">", "lessThan": "<", "greaterThanOrEqual": ">=", "lessThanOrEqual": "<="}
    vals = [formula_display(f) for f in rule.formulas]
    if k == "cellIs":
        if rule.op in ("between", "notBetween") and len(vals) >= 2:
            return f"Cell value {ops[rule.op]} {vals[0]} and {vals[1]}"
        return f"Cell value {ops.get(rule.op, rule.op)} {vals[0] if vals else ''}"
    if k == "expression":
        return f"Formula: ={rule.formulas[0].lstrip('=') if rule.formulas else ''}"
    if k in TEXT_KINDS:
        return {"containsText": "Text contains", "notContainsText": "Text does not contain",
                "beginsWith": "Text begins with", "endsWith": "Text ends with"}[k] + f' "{rule.text}"'
    if k == "top10":
        return f"{'Bottom' if rule.bottom else 'Top'} {rule.rank}{'%' if rule.percent else ''}"
    if k == "aboveAverage":
        return "Above average" if rule.above else "Below average"
    return {"duplicateValues": "Duplicate values", "uniqueValues": "Unique values",
            "containsBlanks": "Blank cells", "notContainsBlanks": "No blanks",
            "containsErrors": "Cells with errors", "notContainsErrors": "No errors",
            "colorScale": "Color scale", "dataBar": "Data bar"}.get(k, k)
